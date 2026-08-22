# -*- coding: utf-8 -*-
"""
pipeline.py —— 两阶段转换核心：图片 → 像素画 → Minecraft 投影 (.litematic)

阶段1（像素画）：
    把任意图片缩放到指定像素尺寸（16×16 ~ 4096×4096 或自定义宽高，可选保持宽高比），
    每格 = 1 像素，保留原图颜色（不做任何色板限制）。可选去除背景（不规则裁剪）。

阶段2（MC 投影）：
    把像素画逐格匹配到最近的 Minecraft 方块颜色，写出 .litematic 建筑投影
    （Version 7 / SubVersion 1），同时生成方块化预览图与材料清单 Excel。

本模块无 GUI 依赖，可供图形界面 main.py 与命令行 cli.py 共同调用。
"""
from __future__ import annotations

import os
from collections import Counter

from PIL import Image, ImageDraw

import pixelart2litematic
from mc_palette import ColorMatcher, get_palette, GLASS_IDS

# 尺寸预设（多比例）
SIZE_PRESETS = ["16×16", "24×24", "32×32", "48×48", "64×64", "96×96",
                "128×128", "256×256", "512×512", "1024×1024", "2048×2048",
                "4096×4096", "自定义"]
MAX_DIM = 4096

# MC 侧调色板选项
MC_PALETTE_KINDS = [
    "auto（全部149种）", "wool（羊毛）", "concrete（混凝土）",
    "terracotta（陶瓦）", "glass（玻璃）", "wool+concrete",
    "concrete+terracotta", "misc（建材）",
]
MC_PALETTE_VALUES = ["auto", "wool", "concrete", "terracotta", "glass",
                     "wool+concrete", "concrete+terracotta", "misc"]
DITHER_KINDS = ["floyd（抖动，渐变平滑）", "none（纯色）"]
DITHER_VALUES = ["floyd", "none"]


# ---------------------------------------------------------------------------
# 阶段1：图片 → 像素画网格（保留原色）
# ---------------------------------------------------------------------------
def load_image_rgb(path: str) -> Image.Image:
    """读取图片并统一为 RGB（透明区域合成到白底）。"""
    img = Image.open(path)
    img.load()
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
    if img.mode == "RGBA":
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[-1])
        img = bg
    elif img.mode != "RGB":
        img = img.convert("RGB")
    return img


def compute_target_size(src_w, src_h, target_w, target_h, preserve_aspect=False):
    """计算实际像素画尺寸。

    preserve_aspect=False：直接使用目标宽高（拉伸）。
    preserve_aspect=True ：保持原图宽高比，按目标框等比缩放（较长边 = 目标值）。
    """
    tw, th = max(1, int(target_w)), max(1, int(target_h))
    if not preserve_aspect:
        return tw, th
    scale = min(tw / src_w, th / src_h)
    w = max(1, round(src_w * scale))
    h = max(1, round(src_h * scale))
    return w, h


def resize_to_grid(image_rgb, width, height, preserve_aspect=False, progress_cb=None):
    """把图片缩放到目标像素尺寸，返回 list[list[tuple|None]]，每个元素为 (r,g,b)。"""
    w, h = compute_target_size(image_rgb.width, image_rgb.height,
                               width, height, preserve_aspect)
    resized = image_rgb.resize((w, h), Image.LANCZOS)
    px = resized.load()
    grid = []
    for y in range(h):
        row = []
        for x in range(w):
            row.append(tuple(px[x, y][:3]))
        grid.append(row)
        if progress_cb is not None and (y % 64 == 0 or y == h - 1):
            progress_cb(y + 1, h)
    return grid


# ---------------------------------------------------------------------------
# 背景去除（不规则裁剪，基于原始像素）
# ---------------------------------------------------------------------------
def _redmean_distance(c1, c2):
    """带红均值加权的 RGB 距离（更接近人眼感知）。"""
    r1, g1, b1 = c1
    r2, g2, b2 = c2
    rmean = (r1 + r2) / 2.0
    dr, dg, db = r1 - r2, g1 - g2, b1 - b2
    return ((2 + rmean / 256.0) * dr * dr + 4.0 * dg * dg
            + (2 + (255 - rmean) / 256.0) * db * db) ** 0.5


def _is_white_pixel(rgb, min_val=195, max_spread=45):
    """判断是否属于“近白色”（亮度高且饱和度低），用于识别空白底色。"""
    r, g, b = rgb
    return min(r, g, b) >= min_val and (max(r, g, b) - min(r, g, b)) <= max_spread


def _background_mask(grid, bg_color=None, tolerance=40):
    """返回与图像边缘相连的背景格子掩码（4 连通洪泛，True=背景）。

    :param bg_color: 背景 RGB 三元组；None 时使用白色启发式（近白色）。
    返回: (mask, removed)；mask 为 list[list[bool]]，removed 为背景格子数。
    """
    if bg_color is not None:
        bg = tuple(bg_color)

        def matches(c):
            return _redmean_distance(c, bg) <= tolerance
    else:
        matches = _is_white_pixel

    h = len(grid)
    if h == 0:
        return [], 0
    w = len(grid[0])

    bg_match = [[matches(grid[y][x]) for x in range(w)] for y in range(h)]
    mask = [[False] * w for _ in range(h)]
    stack = []
    for x in range(w):
        if bg_match[0][x]:
            stack.append((0, x))
        if bg_match[h - 1][x]:
            stack.append((h - 1, x))
    for y in range(h):
        if bg_match[y][0]:
            stack.append((y, 0))
        if bg_match[y][w - 1]:
            stack.append((y, w - 1))
    while stack:
        y, x = stack.pop()
        if mask[y][x] or not bg_match[y][x]:
            continue
        mask[y][x] = True
        for dy, dx in ((1, 0), (-1, 0), (0, 1), (0, -1)):
            ny, nx = y + dy, x + dx
            if 0 <= ny < h and 0 <= nx < w and not mask[ny][nx] and bg_match[ny][nx]:
                stack.append((ny, nx))
    return mask, sum(sum(row) for row in mask)


def remove_background(grid, bg_color=None, tolerance=40):
    """去除“主体之外”的背景，保留主体形状（不规则裁剪）。

    与 _background_mask 同一套背景判定；此处把背景格子置为 None 并裁剪到
    主体最小包围盒。返回: (new_grid, removed, bbox)；bbox = (top, bottom, left, right)。
    """
    h = len(grid)
    if h == 0:
        return grid, 0, (0, 0, 0, 0)
    w = len(grid[0])
    mask, removed = _background_mask(grid, bg_color=bg_color, tolerance=tolerance)
    if removed == 0 or removed == h * w:
        return grid, 0, (0, 0, 0, 0)   # 无背景 / 整幅图全为背景

    ys = [y for y in range(h) for x in range(w) if not mask[y][x]]
    xs = [x for y in range(h) for x in range(w) if not mask[y][x]]
    top, bottom = min(ys), max(ys) + 1
    left, right = min(xs), max(xs) + 1

    new_grid = []
    for y in range(top, bottom):
        row = []
        for x in range(left, right):
            row.append(None if mask[y][x] else grid[y][x])
        new_grid.append(row)
    return new_grid, removed, (top, bottom, left, right)


def fill_background_glass(block_grid, color_grid, mask, color_space="lab"):
    """把背景掩码格子强制替换为最接近其原色的玻璃方块。

    用于"背景填玻璃"模式：白色/取样色背景 → 整片玻璃，主体保持实体方块。
    返回新的 block_grid（原地修改）。
    """
    h, w = len(block_grid), len(block_grid[0])
    cells = [(y, x) for y in range(h) for x in range(w) if mask[y][x]]
    if not cells:
        return block_grid
    matcher = ColorMatcher(get_palette("glass"), color_space)
    rgbs = []
    keep = []
    for y, x in cells:
        c = color_grid[y][x]
        if c is None:
            continue
        rgbs.append(tuple(c))
        keep.append((y, x))
    if not rgbs:
        return block_grid
    if matcher.vectorized:
        try:
            import numpy as np
            arr = np.asarray(rgbs, dtype=np.float64)
            ids, rgb_arr = matcher.match_batch(arr)
            for (y, x), gid, grgb in zip(keep, ids, rgb_arr):
                block_grid[y][x] = (gid, (int(grgb[0]), int(grgb[1]), int(grgb[2])))
            return block_grid
        except ImportError:
            pass
    for (y, x), c in zip(keep, rgbs):
        block_grid[y][x] = matcher.match(c)
    return block_grid


def grid_to_rgba_image(grid):
    """把像素网格（list[list[tuple|None]]）转成 RGBA 图片。

    None（背景已去除）→ 透明像素；实心格 → 原色。
    阶段2 用这张内存图做 MC 方块匹配，透明像素会变成空气（镂空效果）。
    """
    h = len(grid)
    w = len(grid[0]) if h else 0
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    px = img.load()
    for y in range(h):
        for x in range(w):
            c = grid[y][x]
            if c is not None:
                px[x, y] = c + (255,)
    return img


def render_grid_preview(grid, cell=16, show_grid=False, empty_fill=(243, 243, 243)):
    """把像素网格渲染成预览图（空格子用浅灰填充，可叠加网格线）。"""
    h, w = len(grid), len(grid[0])
    img = Image.new("RGB", (w, h), tuple(empty_fill))
    px = img.load()
    for y in range(h):
        for x in range(w):
            c = grid[y][x]
            if c is not None:
                px[x, y] = c
    if cell > 1:
        img = img.resize((w * cell, h * cell), Image.NEAREST)
    if show_grid and cell >= 4:
        draw = ImageDraw.Draw(img)
        for i in range(w + 1):
            draw.line([(i * cell, 0), (i * cell, h * cell)], fill=(40, 40, 40), width=1)
        for j in range(h + 1):
            draw.line([(0, j * cell), (w * cell, j * cell)], fill=(40, 40, 40), width=1)
    return img


def render_grid_pure(grid):
    """渲染 1px/格 的纯像素画（None → 透明，导出用）。"""
    h, w = len(grid), len(grid[0])
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    px = img.load()
    for y in range(h):
        for x in range(w):
            c = grid[y][x]
            if c is not None:
                px[x, y] = c + (255,)
    return img


class Stage1Result:
    """阶段1 输出：像素画网格。"""

    __slots__ = ("grid", "size", "trimmed", "removed", "bbox", "pure_image",
                 "background", "background_mask")

    def __init__(self, grid, size, trimmed=False, removed=0, bbox=None,
                 pure_image=None, background="none", background_mask=None):
        self.grid = grid                  # list[list[tuple|None]]
        self.size = size                  # (w, h)
        self.trimmed = trimmed            # 是否执行了背景处理（去除/填玻璃）
        self.removed = removed            # 识别到的背景格子数
        self.bbox = bbox                  # (top, bottom, left, right) 或 None
        self.pure_image = pure_image      # 1px/格 的纯像素画 PIL Image
        self.background = background      # "none" / "trim" / "glass"
        self.background_mask = background_mask  # 仅 glass 模式：list[list[bool]]


def run_stage1(image_rgb, width, height, *, background="none", bg_color=None,
               preserve_aspect=False, progress_cb=None):
    """阶段1：图片 → 像素画网格（保留原色，多比例尺寸）。

    :param image_rgb: PIL RGB 图片（load_image_rgb 的返回值）
    :param width, height: 目标像素尺寸（preserve_aspect 时作为参考框）
    :param background: "none"=不处理 / "trim"=去除背景（镂空裁剪）/
                       "glass"=背景填玻璃（背景格子 → 玻璃，主体保持实体）
    :param bg_color: 背景 RGB 元组；None=自动白色启发式
    :param preserve_aspect: 保持原图宽高比
    :param progress_cb: 可选回调 progress_cb(pct: float, stage: str)
    """
    def report(pct, stage):
        if progress_cb:
            progress_cb(pct, stage)

    if background not in ("none", "trim", "glass"):
        raise ValueError(f"background 必须为 none/trim/glass，收到: {background}")

    report(5.0, "阶段1 缩放图片")
    w, h = compute_target_size(image_rgb.width, image_rgb.height,
                               width, height, preserve_aspect)

    def p(done, total):
        report(6.0 + 32.0 * done / total, f"阶段1 缩放 {w}×{h}… {done}/{total} 行")

    grid = resize_to_grid(image_rgb, w, h, preserve_aspect=False, progress_cb=p)
    report(40.0, "阶段1 完成缩放")

    trimmed, removed, bbox = False, 0, None
    bg_mask = None
    if background == "trim":
        report(42.0, "阶段1 去除背景")
        grid, removed, bbox = remove_background(grid, bg_color=bg_color)
        trimmed = removed > 0
    elif background == "glass":
        report(42.0, "阶段1 识别背景")
        bg_mask, removed = _background_mask(grid, bg_color=bg_color)
        trimmed = removed > 0

    report(47.0, "阶段1 生成像素画")
    pure_image = render_grid_pure(grid)
    report(48.0, "阶段1 完成")
    return Stage1Result(grid=grid, size=(len(grid[0]), len(grid)),
                        trimmed=trimmed, removed=removed, bbox=bbox,
                        pure_image=pure_image, background=background,
                        background_mask=bg_mask)


# ---------------------------------------------------------------------------
# 玻璃一致性去噪
# ---------------------------------------------------------------------------
def _neighbors4(grid, y, x):
    """返回 (y,x) 的四连通邻居坐标列表（限界内）。"""
    h, w = len(grid), len(grid[0])
    out = []
    for dy, dx in ((1, 0), (-1, 0), (0, 1), (0, -1)):
        ny, nx = y + dy, x + dx
        if 0 <= ny < h and 0 <= nx < w:
            out.append((ny, nx))
    return out


def _neighbors8(y, x, h, w):
    """遍历 (y,x) 的八连通邻居坐标（限界内）。"""
    for dy in (-1, 0, 1):
        for dx in (-1, 0, 1):
            if dy == 0 and dx == 0:
                continue
            ny, nx = y + dy, x + dx
            if 0 <= ny < h and 0 <= nx < w:
                yield ny, nx


def _most_frequent(items):
    """返回 (block_id, rgb) 列表中出现次数最多的项；同票时取第一个。"""
    counter = Counter(bid for bid, _ in items)
    best = max(counter.values())
    for it in items:
        if counter[it[0]] == best:
            return it
    return items[0]


def _small_components(mask, h, w, max_area):
    """返回 mask（list[bytearray]，1=属于该类）中面积 <= max_area 的八连通成分。

    大成分会提前中止并整片标记跳过，总体工作量约 O(格子数)。
    """
    state = [bytearray(w) for _ in range(h)]   # 0=未访问 1=大片(不翻) 2=已归入小片
    comps = []
    for y in range(h):
        row_mask = mask[y]
        for x in range(w):
            if not row_mask[x] or state[y][x]:
                continue
            stack = [(y, x)]
            state[y][x] = 2
            comp = []
            large = False
            while stack:
                cy, cx = stack.pop()
                comp.append((cy, cx))
                if len(comp) > max_area:
                    large = True
                    break
                for ny, nx in _neighbors8(cy, cx, h, w):
                    if mask[ny][nx] and state[ny][nx] == 0:
                        state[ny][nx] = 2
                        stack.append((ny, nx))
            if large:
                for cy, cx in comp:
                    state[cy][cx] = 1
            else:
                comps.append(comp)
    return comps


def denoise_glass(block_grid, glass_ids=None, max_comp_area=10, iterations=2,
                  fill_solid_holes=True):
    """基于连通域的玻璃一致性清理：让玻璃只出现在成片区域。

    - 面积不超过 max_comp_area 的**玻璃小块**（嵌在实体方块中）→ 整体换成
      周围最常见的实体方块（消除白色图案中间的散点玻璃，2×2/3×3 簇也能清掉）
    - fill_solid_holes=True 时：面积不超过 max_comp_area 的**实体小块**
      （嵌在玻璃中）→ 整体补成周围最常见的玻璃（背景玻璃区域保持均匀）
    - 成片大区域（面积 > max_comp_area）不受影响；空气（镂空）永不改动。

    「背景填玻璃」模式下 fill_solid_holes 应传 False——此时玻璃是整片背景，
    打开该规则会把玻璃背景中面积较小的主体碎块误翻成玻璃。

    返回 (new_grid, removed_glass, added_glass)。
    """
    glass_ids = set(glass_ids) if glass_ids is not None else set(GLASS_IDS)
    h, w = len(block_grid), len(block_grid[0])
    grid = [row[:] for row in block_grid]
    removed = added = 0
    for _ in range(iterations):
        is_glass = [bytearray(1 if grid[y][x][0] in glass_ids else 0 for x in range(w))
                    for y in range(h)]
        is_solid = [bytearray(1 if (not is_glass[y][x]) and grid[y][x][0] != "minecraft:air"
                              else 0 for x in range(w))
                    for y in range(h)]
        changed = 0

        # 1) 小片玻璃 → 周围最常见的实体
        for comp in _small_components(is_glass, h, w, max_comp_area):
            comp_set = set(comp)
            border = []
            for cy, cx in comp:
                for ny, nx in _neighbors8(cy, cx, h, w):
                    if (ny, nx) not in comp_set:
                        border.append((ny, nx))
            solids = [b for b in border if is_solid[b[0]][b[1]]]
            airs = [b for b in border if not is_solid[b[0]][b[1]]]
            if solids and len(solids) >= len(airs):
                best = _most_frequent([grid[ny][nx] for ny, nx in solids])
                for cy, cx in comp:
                    grid[cy][cx] = best
                changed += len(comp)
                removed += len(comp)

        # 2) 小片实体 → 周围最常见的玻璃（fill_solid_holes=False 时跳过）
        if fill_solid_holes:
            for comp in _small_components(is_solid, h, w, max_comp_area):
                comp_set = set(comp)
                border = []
                for cy, cx in comp:
                    for ny, nx in _neighbors8(cy, cx, h, w):
                        if (ny, nx) not in comp_set:
                            border.append((ny, nx))
                glasses = [b for b in border if is_glass[b[0]][b[1]]]
                airs = [b for b in border if not is_glass[b[0]][b[1]]]
                if glasses and len(glasses) >= len(airs):
                    best = _most_frequent([grid[ny][nx] for ny, nx in glasses])
                    for cy, cx in comp:
                        grid[cy][cx] = best
                    changed += len(comp)
                    added += len(comp)

        if changed == 0:
            break
    return grid, removed, added


# ---------------------------------------------------------------------------
# 阶段2：像素画网格 → MC 方块 → .litematic
# ---------------------------------------------------------------------------
def run_stage2(stage1: Stage1Result, *, mc_palette="auto", dither="floyd",
               color_space="lab", scale=1, thickness=1, backing=None,
               origin=(0, 0, 0), name="像素画", author="", description="",
               mc_data_version=3955, out_path=None, preview=True, stats=True,
               cell=14, orientation="wall", glass_cleanup=True,
               progress_cb=None):
    """阶段2：把像素画网格转成 .litematic 投影。

    透明格子（None，背景已去除）→ 空气，形成镂空效果。

    :param orientation: 'wall'=竖放（墙面/平视，图像在 X-Y 平面）；
                        'floor'=横放（地面/俯视，图像在 X-Z 平面）。
    :param glass_cleanup: 消除孤立玻璃/实体噪点，让玻璃只出现在成片区域
                          （调色板不含玻璃时自动跳过）。

    返回结果字典：
        {out, preview, stats, w, h, total_blocks, distinct_blocks, raw_size,
         file_size, block_grid, mc_palette_kind, color_space, scale, thickness,
         orientation, glass_fixed}
    """
    def report(pct, stage):
        if progress_cb:
            progress_cb(pct, stage)

    report(50.0, "阶段2 匹配方块颜色")
    img = grid_to_rgba_image(stage1.grid)
    has_glass = bool(set(get_palette(mc_palette)) & set(GLASS_IDS))
    # 玻璃门槛：玻璃必须明显优于最佳实体方块才选用（避免白色区域被玻璃"恰好命中"）
    matcher = ColorMatcher(get_palette(mc_palette), color_space,
                           glass_bias=0.75 if (glass_cleanup and has_glass) else None,
                           glass_margin=100)
    block_grid, (w, h) = pixelart2litematic.match_image(
        img, matcher, dither=dither,
        progress_cb=lambda p, s: report(50.0 + p * 0.35, "阶段2 匹配方块颜色"))

    glass_fixed = (0, 0)
    bg_glass_cells = 0
    if has_glass:
        bg_mask = stage1.background_mask
        if bg_mask is not None:
            # "背景填玻璃"模式：背景格子强制替换为最接近其原色的玻璃
            report(86.0, "阶段2 背景填充玻璃")
            block_grid = fill_background_glass(block_grid, stage1.grid, bg_mask,
                                               color_space)
            bg_glass_cells = sum(1 for y in range(h) for x in range(w)
                                 if bg_mask[y][x] and block_grid[y][x][0] in GLASS_IDS)
        if glass_cleanup:
            report(87.0, "阶段2 玻璃去噪")
            block_grid, rem, add = denoise_glass(
                block_grid, fill_solid_holes=(bg_mask is None))
            glass_fixed = (rem, add)

    report(88.0, "阶段2 生成投影")
    sch, counts = pixelart2litematic.build_schematic(
        block_grid, (w, h), tuple(origin), scale, thickness, backing,
        name, author, description, mc_data_version,
        progress_cb=lambda done, total: report(
            88.0 + 6.0 * done / total, f"阶段2 生成投影… {done}/{total} 行"),
        orientation=orientation)

    if out_path is None:
        out_path = name + ".litematic"
    raw_size = sch.save(out_path)
    total_blocks = sum(counts.values())

    preview_path = stats_path = None
    if preview:
        preview_path = os.path.splitext(out_path)[0] + ".preview.png"
        report(95.0, "阶段2 渲染预览")
        pixelart2litematic.render_preview(block_grid, (w, h), cell=cell,
                                          out_path=preview_path)
    if stats:
        stats_path = os.path.splitext(out_path)[0] + ".stats.xlsx"
        report(96.5, "阶段2 生成材料清单")
        stats_path = write_stats_xlsx(stats_path, counts)
    report(100.0, "完成")

    return {
        "out": out_path,
        "preview": preview_path,
        "stats": stats_path,
        "w": w, "h": h,
        "total_blocks": total_blocks,
        "distinct_blocks": len(counts),
        "raw_size": raw_size,
        "file_size": os.path.getsize(out_path),
        "block_grid": block_grid,
        "mc_palette_kind": mc_palette,
        "color_space": color_space,
        "scale": scale,
        "thickness": thickness,
        "orientation": orientation,
        "glass_fixed": glass_fixed,
        "bg_glass_cells": bg_glass_cells,
    }


# ---------------------------------------------------------------------------
# 完整管道
# ---------------------------------------------------------------------------
def run_pipeline(image_path, *, width=48, height=48, preserve_aspect=False,
                 background="none", bg_color=None,
                 mc_palette="auto", dither="floyd", color_space="lab",
                 scale=1, thickness=1, backing=None, origin=(0, 0, 0),
                 name="像素画", author="", description="", mc_data_version=3955,
                 out_dir=None, mc_preview=True, mc_stats=True, cell=14,
                 orientation="wall", glass_cleanup=True,
                 export_pixel_art=False, progress_cb=None):
    """完整两阶段转换：图片 → 像素画 → .litematic。

    :param background: "none"=不处理 / "trim"=去除背景（镂空）/
                       "glass"=背景填玻璃（白色/取样色背景 → 玻璃，主体保持实体）

    输出文件（默认与图片同目录，文件名同图片名）：
        - {stem}.litematic            MC 建筑投影（游戏内加载）
        - {stem}.preview.png          方块化预览图
        - {stem}.stats.xlsx          方块材料清单（中文名/数量/色块）
        - {stem}_像素画.png            1px/格 纯像素画（可选）

    返回 dict：{stage1: Stage1Result, stage2: dict, image_path, stem}
    """
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"找不到图片: {image_path}")

    def report(pct, stage):
        if progress_cb:
            progress_cb(pct, stage)

    report(1.0, "读取图片")
    image_rgb = load_image_rgb(image_path)
    stem = os.path.splitext(os.path.basename(image_path))[0]
    base_dir = out_dir or os.path.dirname(os.path.abspath(image_path))
    os.makedirs(base_dir, exist_ok=True)

    stage1 = run_stage1(
        image_rgb, width, height,
        background=background, bg_color=bg_color,
        preserve_aspect=preserve_aspect, progress_cb=report,
    )

    out_path = os.path.join(base_dir, stem + ".litematic")
    stage2 = run_stage2(
        stage1, mc_palette=mc_palette, dither=dither, color_space=color_space,
        scale=scale, thickness=thickness, backing=backing, origin=origin,
        name=name, author=author, description=description,
        mc_data_version=mc_data_version, out_path=out_path,
        preview=mc_preview, stats=mc_stats, cell=cell,
        orientation=orientation, glass_cleanup=glass_cleanup,
        progress_cb=report,
    )

    pixel_art_path = None
    if export_pixel_art:
        pixel_art_path = os.path.join(base_dir, stem + "_像素画.png")
        stage1.pure_image.save(pixel_art_path)

    return {
        "stage1": stage1,
        "stage2": stage2,
        "image_path": image_path,
        "stem": stem,
        "pixel_art_path": pixel_art_path,
    }


# ---------------------------------------------------------------------------
# 材料清单导出（Excel）
# ---------------------------------------------------------------------------
def write_stats_xlsx(path, counts):
    """把方块用量写为 Excel 材料清单（.xlsx）。

    列：序号 | 方块(中文) | 方块ID | 数量 | 颜色HEX | 颜色预览(色块)。
    未安装 openpyxl 时回退为 CSV（扩展名 .csv），保证始终能导出。
    返回实际写入的文件路径。
    """
    import mc_names
    from mc_palette import ALL_BLOCKS

    def _csv_fallback():
        csv_path = os.path.splitext(path)[0] + ".csv"
        pixelart2litematic.write_stats(csv_path, counts)
        return csv_path

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return _csv_fallback()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "材料清单"

    total = sum(counts.values())
    items = sorted(counts.items(), key=lambda kv: -kv[1])

    # 标题 + 信息
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "像素画方块材料清单"
    c.font = Font(size=14, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    ws["A2"] = f"共 {len(items)} 种方块，合计 {total} 个"
    ws["A2"].font = Font(size=10, color="666666")

    # 表头
    headers = ["序号", "方块（中文）", "方块ID", "数量", "颜色HEX", "颜色预览"]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 数据行
    for i, (bid, cnt) in enumerate(items, start=1):
        row = 3 + i
        rgb = ALL_BLOCKS.get(bid, (0, 0, 0))
        hex_str = "#%02X%02X%02X" % rgb
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=mc_names.cn_name(bid))
        ws.cell(row=row, column=3, value=bid.replace("minecraft:", ""))
        ws.cell(row=row, column=4, value=cnt)
        ws.cell(row=row, column=5, value=hex_str)
        preview = ws.cell(row=row, column=6)
        preview.fill = PatternFill("solid", fgColor=hex_str.lstrip("#"))
        preview.alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=4).alignment = Alignment(horizontal="right")

    # 合计行
    last = 3 + len(items)
    ws.cell(row=last + 1, column=2, value="合计").font = Font(bold=True)
    tot = ws.cell(row=last + 1, column=4, value=total)
    tot.font = Font(bold=True)
    tot.alignment = Alignment(horizontal="right")

    # 列宽 + 冻结表头
    widths = [6, 22, 24, 10, 12, 12]
    for col, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = wd
    ws.freeze_panes = "A4"

    try:
        wb.save(path)
        return path
    except Exception:
        return _csv_fallback()


# ---------------------------------------------------------------------------
# 便捷函数
# ---------------------------------------------------------------------------
def block_usage_summary(stage2):
    """返回 MC 方块用量的可读摘要行。"""
    r = stage2
    orient = "竖放(墙面/平视)" if r.get("orientation", "wall") == "wall" else "横放(地面/俯视)"
    return (f"{r['w']}×{r['h']} 像素 → {r['w'] * r['scale']}×{r['h'] * r['scale']}"
            f"×{r['thickness']} 方块，共 {r['total_blocks']} 个 / "
            f"{r['distinct_blocks']} 种，"
            f"调色板 {r['mc_palette_kind']}，颜色空间 {r['color_space']}，{orient}")
