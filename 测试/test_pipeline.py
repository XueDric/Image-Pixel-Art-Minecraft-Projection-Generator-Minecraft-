# -*- coding: utf-8 -*-
"""整合管道测试：
1. 端到端两阶段转换（多种参数组合）
2. litemapy（独立实现）逐方块比对
3. 内置 NBT 读取器回环解析
4. 镂空（背景去除→空气）验证
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import litemapy
from PIL import Image, ImageDraw

import pipeline
from nbt_dump import load as my_nbt_load

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(HERE, "work")
os.makedirs(WORK, exist_ok=True)


# ---------------------------------------------------------------------------
# 测试图
# ---------------------------------------------------------------------------
def make_gradient(path, w=96, h=72):
    img = Image.new("RGB", (w, h))
    px = img.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = (int(255 * x / w), int(160 * y / h), 120)
    d = ImageDraw.Draw(img)
    d.ellipse([w * 0.2, h * 0.15, w * 0.62, h * 0.8], fill=(235, 90, 70))
    d.rectangle([w * 0.66, h * 0.1, w * 0.95, h * 0.4], fill=(30, 120, 210))
    img.save(path)
    return path


def make_2x2(path):
    """2×2 纯色图：上排 红/白，下排 绿/蓝（绿=0,255,0 取 lime 等最近色）。"""
    img = Image.new("RGB", (2, 2))
    img.putpixel((0, 0), (255, 0, 0))     # 顶行左：红
    img.putpixel((1, 0), (255, 255, 255))  # 顶行右：白
    img.putpixel((0, 1), (0, 255, 0))     # 底行左：绿
    img.putpixel((1, 1), (0, 0, 255))     # 底行右：蓝
    img.save(path)
    return path


def make_trim(path, w=120, h=90):
    """白色背景 + 中央红色圆（用于去除背景测试）。"""
    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.ellipse([w * 0.3, h * 0.2, w * 0.7, h * 0.8], fill=(235, 90, 70))
    img.save(path)
    return path


# ---------------------------------------------------------------------------
# litematic 读取辅助
# ---------------------------------------------------------------------------
def read_blocks(path):
    """用 litemapy 读回方块: 返回 ((x,y,z) -> block_name) 字典 + region。"""
    sch = litemapy.Schematic.load(path)
    reg = list(sch.regions.values())[0]
    out = {}
    for x in range(reg.width):
        for y in range(reg.height):
            for z in range(reg.length):
                b = str(reg[x, y, z]).split("[")[0]
                if b != "minecraft:air":
                    out[(x, y, z)] = b
    return out, reg


# ---------------------------------------------------------------------------
# 用例
# ---------------------------------------------------------------------------
def test_end_to_end_gradient():
    src = make_gradient(os.path.join(WORK, "grad.png"))
    res = pipeline.run_pipeline(
        src, width=64, height=64, mc_palette="auto", dither="floyd",
        color_space="lab", scale=1, thickness=1, out_dir=WORK)
    s1, s2 = res["stage1"], res["stage2"]
    assert s1.size == (64, 64), s1.size
    assert s2["w"] == 64 and s2["h"] == 64
    assert os.path.exists(s2["out"]) and s2["file_size"] > 0
    assert s2["total_blocks"] > 0
    # 阶段1输出：纯像素画为 64×64，保留原色
    assert s1.pure_image.size == (64, 64)
    assert s1.grid[0][0] == s1.grid[0][0][:3] or isinstance(s1.grid[0][0], tuple)
    print(f"  ✅ 端到端渐变图: {s1.size} 格 → {s2['total_blocks']} 方块 / "
          f"{s2['distinct_blocks']} 种 ({s2['file_size']/1024:.1f} KB)")
    return res


def test_2x2_exact_blocks():
    """rgb 空间 + wool 调色板 → 可精确预测方块。"""
    src = make_2x2(os.path.join(WORK, "t2x2.png"))
    res = pipeline.run_pipeline(
        src, width=2, height=2,
        mc_palette="wool", dither="none", color_space="rgb",
        out_dir=WORK)
    s2 = res["stage2"]
    blocks, reg = read_blocks(s2["out"])
    # 图像行0(顶) → 最高 y=1；图像行1(底) → y=0
    expect = {
        (0, 1, 0): "minecraft:red_wool",      # 顶行左 红
        (1, 1, 0): "minecraft:white_wool",    # 顶行右 白
        (0, 0, 0): "minecraft:lime_wool",     # 底行左 绿 → 最近 lime
        (1, 0, 0): "minecraft:blue_wool",     # 底行右 蓝
    }
    assert blocks == expect, f"方块不符:\n 实际 {blocks}\n 期望 {expect}"
    print("  ✅ 2×2 精确方块比对通过:", blocks)


def test_keep_aspect_and_3d():
    """保持宽高比 + 自定义尺寸 + 厚度2 + 背面填充。"""
    src = make_gradient(os.path.join(WORK, "grad2.png"), 160, 120)
    res = pipeline.run_pipeline(
        src, width=64, height=64, preserve_aspect=True,
        mc_palette="concrete", dither="none", color_space="rgb",
        thickness=2, backing="minecraft:white_concrete",
        origin=(100, 64, 200), out_dir=WORK)
    s1, s2 = res["stage1"], res["stage2"]
    # 160:120 → 64:48
    assert s1.size == (64, 48), s1.size
    blocks, reg = read_blocks(s2["out"])
    assert (reg.width, reg.height, reg.length) == (64, 48, 2)
    assert (reg.x, reg.y, reg.z) == (100, 64, 200)
    # 所有前景方块 z=1 处都有白混凝土背面
    for (x, y, z), bid in blocks.items():
        if z == 0:
            assert (x, y, 1) in blocks and blocks[(x, y, 1)] == "minecraft:white_concrete", \
                (x, y)
    print(f"  ✅ 保持宽高比+3D: 尺寸 {s1.size}，背面填充验证通过 "
          f"({len(blocks)} 方块)")
    return res


def test_trim_creates_air():
    """去除背景后，镂空区域在投影中应为空气。"""
    src = make_trim(os.path.join(WORK, "trim.png"))
    res = pipeline.run_pipeline(
        src, width=64, height=64, background="trim",
        mc_palette="wool", dither="none", color_space="rgb", out_dir=WORK)
    s1, s2 = res["stage1"], res["stage2"]
    assert s1.trimmed and s1.removed > 0, (s1.removed, s1.size)
    # 裁剪后网格应小于 64×64（红色圆在 64 格内占不满四角）
    gw, gh = s1.size
    assert gw < 64 and gh < 64, s1.size
    blocks, reg = read_blocks(s2["out"])
    assert (reg.width, reg.height) == (gw, gh), (reg.width, reg.height, gw, gh)
    # 区域中心附近应有实心方块（圆形主体）
    cx, cy = gw // 2, gh // 2
    assert (cx, cy, 0) in blocks, "中心应有方块"
    # 角落应为空气：直接检查区域 (0,0,0) 是否为空气
    sch = litemapy.Schematic.load(s2["out"])
    reg0 = list(sch.regions.values())[0]
    assert str(reg0[0, 0, 0]) == "minecraft:air", "角落应为空气"
    print(f"  ✅ 镂空验证: 裁剪到 {gw}×{gh}，角落=空气，中心={blocks.get((cx, cy, 0))}")
    return res


def test_sampled_bg_removal():
    """指定背景色（非白色）去除验证。"""
    src = os.path.join(WORK, "sampled.png")
    img = Image.new("RGB", (90, 60), (0, 140, 200))     # 蓝色背景
    d = ImageDraw.Draw(img)
    d.ellipse([20, 10, 70, 50], fill=(255, 200, 60))    # 黄色主体
    img.save(src)
    res = pipeline.run_pipeline(
        src, width=48, height=48, background="trim",
        bg_color=(0, 140, 200), mc_palette="auto", out_dir=WORK)
    s1 = res["stage1"]
    assert s1.trimmed and s1.removed > 0, (s1.removed, s1.size)
    gw, gh = s1.size
    assert gw < 48 and gh < 48, s1.size
    print(f"  ✅ 指定背景色去除: 蓝色背景 removed={s1.removed} 格，裁剪到 {gw}×{gh}")


def test_floor_orientation():
    """横放（地面/俯视）：图像在 X-Z 平面，厚度沿 +Y 向上，行0→最大Z。"""
    src = make_2x2(os.path.join(WORK, "t2x2_floor.png"))
    res = pipeline.run_pipeline(
        src, width=2, height=2,
        mc_palette="wool", dither="none", color_space="rgb",
        origin=(10, 20, 30), orientation="floor", out_dir=WORK)
    s2 = res["stage2"]
    assert s2["orientation"] == "floor"
    blocks, reg = read_blocks(s2["out"])
    # 区域：宽=2(x)，高=1(厚度)，长=2(z)；位置=(10,20,30)
    assert (reg.width, reg.height, reg.length) == (2, 1, 2), (reg.width, reg.height, reg.length)
    assert (reg.x, reg.y, reg.z) == (10, 20, 30)
    # read_blocks 返回的是区域局部坐标 (0,0,0) 起；图像行0(顶) → 最大局部 z
    expect = {
        (0, 0, 1): "minecraft:red_wool",    # 顶行左 红 → z=1
        (1, 0, 1): "minecraft:white_wool",  # 顶行右 白
        (0, 0, 0): "minecraft:lime_wool",   # 底行左 绿 → z=0
        (1, 0, 0): "minecraft:blue_wool",   # 底行右 蓝
    }
    assert blocks == expect, f"横放方块不符:\n 实际 {blocks}\n 期望 {expect}"

    # 厚度2 + 背面填充：y+1 层应为背面方块
    res2 = pipeline.run_pipeline(
        src, width=2, height=2,
        mc_palette="wool", dither="none", color_space="rgb",
        origin=(10, 20, 30), orientation="floor", thickness=2,
        backing="minecraft:white_concrete", out_dir=WORK)
    blocks2, reg2 = read_blocks(res2["stage2"]["out"])
    assert (reg2.width, reg2.height, reg2.length) == (2, 2, 2)
    assert blocks2[(0, 1, 1)] == "minecraft:white_concrete", blocks2
    assert blocks2[(1, 1, 0)] == "minecraft:white_concrete"
    print("  ✅ 横放(俯视)朝向: X-Z 平面映射、厚度沿+Y、背面填充验证通过")
    return res


def test_nbt_roundtrip():
    """内置 NBT 读取器回环解析生成的 .litematic。"""
    src = make_gradient(os.path.join(WORK, "grad3.png"))
    res = pipeline.run_pipeline(
        src, width=32, height=32, mc_palette="auto", dither="floyd",
        out_dir=WORK)
    s2 = res["stage2"]
    name, root = my_nbt_load(s2["out"])
    assert root["Version"] == 7 and root["SubVersion"] == 1
    reg = list(root["Regions"].values())[0]
    for key in ("Position", "Size", "BlockStatePalette", "BlockStates",
                "Entities", "TileEntities", "PendingBlockTicks", "PendingFluidTicks"):
        assert key in reg, f"缺少 {key}"
    assert reg["BlockStatePalette"][0]["Name"] == "minecraft:air"
    assert reg["Size"]["x"] == 32 and reg["Size"]["y"] == 32
    assert len(reg["BlockStates"]) > 0
    print(f"  ✅ NBT 回环解析通过 (palette={len(reg['BlockStatePalette'])}, "
          f"BlockStates={len(reg['BlockStates'])} longs)")


def test_unstable_blocks_filtered():
    """不适宜方块应从调色板剔除；蓝冰/受重力方块按需求保留。"""
    from mc_palette import get_palette, UNSTABLE_BLOCKS
    for kind in ("auto", "misc"):
        pal = get_palette(kind)
        for bid in UNSTABLE_BLOCKS:
            assert bid not in pal, f"{kind} 调色板不应包含 {bid}"
    # 保留的方块不受影响：蓝冰（不融化）、受重力方块、雪块、混凝土等
    keep = get_palette("auto")
    for bid in ("minecraft:blue_ice", "minecraft:sand", "minecraft:red_sand",
                "minecraft:gravel", "minecraft:white_concrete_powder",
                "minecraft:snow_block", "minecraft:packed_mud",
                "minecraft:white_concrete", "minecraft:white_stained_glass",
                "minecraft:soul_soil"):
        assert bid in keep, f"应保留 {bid}"
    print(f"  ✅ 不稳定方块已剔除（{len(UNSTABLE_BLOCKS)} 种），auto 调色板现为 "
          f"{len(get_palette('auto'))} 种（蓝冰/沙子/沙砾等已保留）")


def test_denoise_glass_unit():
    """玻璃去噪单元测试：小块玻璃/实体翻转，成片保留，空气不动。"""
    from pipeline import denoise_glass
    W = ("minecraft:white_wool", (233, 236, 236))
    G = ("minecraft:white_stained_glass", (255, 255, 255))
    R = ("minecraft:red_wool", (160, 39, 34))

    # 1) 实体中的孤立玻璃 → 换成周围实体
    grid = [[W] * 5 for _ in range(5)]
    grid[2][2] = G
    out, rem, add = denoise_glass([r[:] for r in grid])
    assert out[2][2][0] == W[0], out[2][2]
    assert rem == 1 and add == 0, (rem, add)

    # 2) 玻璃中的孤立实体 → 换成周围玻璃
    grid = [[G] * 5 for _ in range(5)]
    grid[2][2] = W
    out, rem, add = denoise_glass([r[:] for r in grid])
    assert out[2][2][0] == G[0], out[2][2]
    assert add == 1 and rem == 0, (rem, add)

    # 3) 实体中的 2×2 玻璃簇（4 格 ≤ 阈值）→ 整体清除
    grid = [[W] * 5 for _ in range(5)]
    for y in (1, 2):
        for x in (1, 2):
            grid[y][x] = G
    out, rem, add = denoise_glass([r[:] for r in grid])
    assert all(out[y][x][0] != G[0] for y in range(5) for x in range(5)), "2×2 玻璃簇应清除"
    assert rem == 4, rem

    # 4) 3×3 玻璃块（9 格 ≤ 阈值）→ 整体清除；4×4（16 格 > 阈值）→ 保留成片
    grid = [[W] * 8 for _ in range(8)]
    for y in (2, 3, 4):
        for x in (2, 3, 4):
            grid[y][x] = G
    out, rem, add = denoise_glass([r[:] for r in grid])
    assert rem == 9, rem

    grid = [[W] * 8 for _ in range(8)]
    for y in (2, 3, 4, 5):
        for x in (2, 3, 4, 5):
            grid[y][x] = G
    out, rem, add = denoise_glass([r[:] for r in grid])
    assert sum(1 for row in out for c in row if c[0] == G[0]) == 16, "4×4 成片玻璃应保留"

    # 5) 空气（镂空）永不改动
    grid = [[W] * 3 for _ in range(3)]
    grid[0][0] = ("minecraft:air", (0, 0, 0))
    grid[1][1] = G
    out, _, _ = denoise_glass([r[:] for r in grid])
    assert out[0][0][0] == "minecraft:air"
    print("  ✅ 玻璃去噪单元测试: 孤立/2×2/3×3 清除、4×4 成片保留、空气不动 均通过")


def test_white_no_glass():
    """纯白区域不应出现玻璃（白色应落到雪块/羊毛/混凝土等实体方块）。"""
    from mc_palette import GLASS_IDS
    p = os.path.join(WORK, "white_no_glass.png")
    Image.new("RGB", (16, 16), (255, 255, 255)).save(p)
    res = pipeline.run_pipeline(p, width=16, height=16, mc_palette="auto",
                                dither="none", out_dir=WORK)
    grid = res["stage2"]["block_grid"]
    glass = [bid for row in grid for bid, _ in row if bid in GLASS_IDS]
    assert not glass, f"纯白区域不应有玻璃: {glass[:5]}"
    top = [bid for row in grid for bid, _ in row]
    from collections import Counter
    print(f"  ✅ 纯白无玻璃: 16×16 全白 → {dict(Counter(top).most_common(2))}")

    # 近白（照片常见灰白）同样不应出现玻璃
    p2 = os.path.join(WORK, "nearwhite_no_glass.png")
    img = Image.new("RGB", (32, 32), (244, 246, 244))
    img.save(p2)
    res2 = pipeline.run_pipeline(p2, width=32, height=32, mc_palette="auto",
                                 dither="none", out_dir=WORK)
    grid2 = res2["stage2"]["block_grid"]
    glass2 = [bid for row in grid2 for bid, _ in row if bid in GLASS_IDS]
    assert not glass2, f"近白区域不应有玻璃: {glass2[:5]}"
    print("  ✅ 近白(244,246,244) 无玻璃通过")


def test_glass_consistency_end_to_end():
    """auto 调色板端到端：转换后不允许出现孤立的玻璃（八连通玻璃邻居≥1）。"""
    src = make_gradient(os.path.join(WORK, "glass_grad.png"), 128, 96)
    res = pipeline.run_pipeline(src, width=64, height=64, mc_palette="auto",
                                dither="floyd", out_dir=WORK)
    s2 = res["stage2"]
    grid = s2["block_grid"]
    h, w = len(grid), len(grid[0])
    from pipeline import _neighbors8
    from mc_palette import GLASS_IDS
    glass_count = iso_count = 0
    for y in range(h):
        for x in range(w):
            bid, _ = grid[y][x]
            if bid in GLASS_IDS:
                glass_count += 1
                if not any(grid[ny][nx][0] in GLASS_IDS
                           for ny, nx in _neighbors8(y, x, h, w)):
                    iso_count += 1
    assert iso_count == 0, f"存在 {iso_count} 个孤立玻璃（总数 {glass_count}）"
    print(f"  ✅ 玻璃一致性端到端: 共 {glass_count} 个玻璃，0 个孤立（"
          f"移除 {s2['glass_fixed'][0]} / 填补 {s2['glass_fixed'][1]}）")


def test_glass_background():
    """背景填玻璃：白色背景 → 整片玻璃，主体保持实体（含小块主体不被误翻）。"""
    from mc_palette import GLASS_IDS

    # 场景：白底 + 中央红色圆 + 一个小的独立红色方块（测试主体碎块保护）
    p = os.path.join(WORK, "glass_bg.png")
    img = Image.new("RGB", (120, 120), (252, 252, 250))
    d = ImageDraw.Draw(img)
    d.ellipse([40, 40, 80, 80], fill=(200, 60, 50))
    d.rectangle([92, 92, 99, 99], fill=(200, 60, 50))   # 独立小块主体
    img.save(p)

    res = pipeline.run_pipeline(p, width=48, height=48, background="glass",
                                mc_palette="auto", dither="none", out_dir=WORK)
    s1, s2 = res["stage1"], res["stage2"]
    assert s1.background == "glass" and s1.background_mask is not None
    assert s1.trimmed and s1.removed > 0, (s1.removed,)
    assert s2["bg_glass_cells"] > 0, s2["bg_glass_cells"]

    grid = s2["block_grid"]
    h, w = len(grid), len(grid[0])
    # 四角应为玻璃（背景）
    for y, x in ((0, 0), (0, w - 1), (h - 1, 0), (h - 1, w - 1)):
        assert grid[y][x][0] in GLASS_IDS, f"角落 ({y},{x}) 应为玻璃: {grid[y][x]}"
    # 中心圆（主体）应为实体，不是玻璃
    assert grid[h // 2][w // 2][0] not in GLASS_IDS, "主体中心不应是玻璃"
    # 独立小块主体（右下角附近）应保持实体，不被玻璃吞噬
    small = [(y, x) for y in range(h) for x in range(w)
             if s1.background_mask[y][x] is False
             and not any(s1.background_mask[ny][nx] for ny, nx in
                         ((y + 1, x), (y - 1, x), (y, x + 1), (y, x - 1))
                         if 0 <= ny < h and 0 <= nx < w)]
    solid_small = [c for c in small if grid[c[0]][c[1]][0] not in GLASS_IDS]
    assert solid_small, "应存在保持实体的主体格子"
    print(f"  ✅ 背景填玻璃: 背景 {s1.removed} 格 → 玻璃 {s2['bg_glass_cells']} 格，"
          f"主体中心/独立小块均保持实体")

    # 填玻璃应匹配到白色玻璃
    corners = {grid[0][0][0], grid[0][w - 1][0], grid[h - 1][0][0], grid[h - 1][w - 1][0]}
    print(f"     角落玻璃方块: {corners}")


def test_stats_xlsx_material_list():
    """材料清单应生成 Excel（中文方块名/数量/合计），且行数等于方块种类。"""
    src = make_gradient(os.path.join(WORK, "xlsx_grad.png"), 64, 64)
    res = pipeline.run_pipeline(src, width=32, height=32, mc_palette="auto",
                                dither="floyd", out_dir=WORK)
    s2 = res["stage2"]
    xlsx = s2["stats"]
    assert xlsx and xlsx.endswith(".xlsx") and os.path.exists(xlsx), xlsx

    import openpyxl
    import mc_names
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    assert ws["A1"].value == "像素画方块材料清单", ws["A1"].value
    # 表头
    assert [ws.cell(row=3, column=c).value for c in range(1, 7)] == \
        ["序号", "方块（中文）", "方块ID", "数量", "颜色HEX", "颜色预览"]
    # 数据行 = 方块种类数；每种方块都有中文名且数量>0；合计正确
    distinct = s2["distinct_blocks"]
    rows = [r for r in ws.iter_rows(min_row=4, max_row=3 + distinct, max_col=4)]
    assert len(rows) == distinct, (len(rows), distinct)
    total = 0
    for r in rows:
        cn, bid, cnt = r[1].value, r[2].value, r[3].value
        assert isinstance(cnt, int) and cnt > 0
        assert mc_names.BLOCK_CN.get("minecraft:" + bid), f"缺少中文名: {bid}"
        total += cnt
    assert ws.cell(row=4 + distinct, column=4).value == total == s2["total_blocks"]
    print(f"  ✅ 材料清单 Excel: {distinct} 种方块全部中文化，合计 {total} 个"
          f"（示例: {rows[0][1].value} ×{rows[0][3].value}）")


def test_progress_callback():
    """进度回调应单调递增到 100。"""
    src = make_gradient(os.path.join(WORK, "grad4.png"), 48, 48)
    seen = []
    pipeline.run_pipeline(src, width=16, height=16, out_dir=WORK,
                          progress_cb=lambda p, s: seen.append((p, s)))
    assert seen and seen[-1][0] == 100.0
    assert all(seen[i][0] <= seen[i + 1][0] + 1e-9 for i in range(len(seen) - 1)), \
        "进度应单调递增"
    print(f"  ✅ 进度回调: {len(seen)} 次，末次 {seen[-1]}")


def main():
    import sys as _sys
    for stream in (_sys.stdout, _sys.stderr):
        try:
            if stream is not None and hasattr(stream, "reconfigure"):
                stream.reconfigure(errors="replace")
        except Exception:
            pass
    print("== 整合管道测试 ==")
    test_end_to_end_gradient()
    test_2x2_exact_blocks()
    test_keep_aspect_and_3d()
    test_floor_orientation()
    test_trim_creates_air()
    test_sampled_bg_removal()
    test_glass_background()
    test_unstable_blocks_filtered()
    test_denoise_glass_unit()
    test_white_no_glass()
    test_glass_consistency_end_to_end()
    test_nbt_roundtrip()
    test_stats_xlsx_material_list()
    test_progress_callback()
    print("\n🎉 全部测试通过：两阶段转换 → litemapy 逐方块比对 → NBT 回环，均一致")


if __name__ == "__main__":
    main()
